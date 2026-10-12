from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

from app.schemas.export_style import ChangeHighlighting
from app.services.export_diff import edition_segments
from app.schemas.protocol_form import ProtocolForm
from app.services.protocol_form import CLOSING_TEXT
from app.services.protocol_generator import edition_label


class ProtocolExporter:
    def export(
        self,
        rows: list[dict],
        drafting_role: str,
        counterparty_role: str,
        contract_name: str | None,
        *,
        form: ProtocolForm | None = None,
        drafting_party_id: str | None = None,
        highlighting: ChangeHighlighting | None = None,
    ) -> bytes:
        form = form or ProtocolForm(
            contract_reference=contract_name or "",
            original_title=edition_label(counterparty_role, "Стороны 1"),
            proposed_title=edition_label(drafting_role, "Стороны 2"),
            opening_text="Стороны договорились заключить договор с учетом настоящего протокола разногласий:",
            closing_text=CLOSING_TEXT,
        )
        doc = DocxDocument()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21)
        section.top_margin = section.bottom_margin = Cm(1.5)
        section.left_margin = section.right_margin = Cm(2)
        normal = doc.styles["Normal"]
        normal.font.name, normal.font.size = "Times New Roman", Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.05

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.keep_with_next = True
        run = title.add_run(form.title)
        run.bold, run.font.size = True, Pt(14)
        reference = doc.add_paragraph("К договору: " + (form.contract_reference or "№ __________ от __________________"))
        reference.alignment = WD_ALIGN_PARAGRAPH.CENTER
        reference.paragraph_format.keep_with_next = True
        dateline = doc.add_paragraph()
        dateline.paragraph_format.tab_stops.add_tab_stop(Cm(25.7), WD_TAB_ALIGNMENT.RIGHT)
        dateline.add_run((form.place or "Место: __________________") + "\t" + (form.date or "Дата: __________________"))
        for text in form.opening_text.splitlines():
            doc.add_paragraph(text)

        table = doc.add_table(rows=1 + len(rows), cols=4)
        table.style = "Table Grid"
        table.autofit = False
        widths = [Cm(2), Cm(7.9), Cm(7.9), Cm(7.9)]
        for column, width in zip(table.columns, widths):
            column.width = width
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        headers = [form.point_title, form.original_title, form.proposed_title, form.agreed_title]
        values = [headers, *[[row[key] for key in ("point_label", "party_1_text", "party_2_text", "agreed_text")] for row in rows]]
        for index, values_row in enumerate(values):
            segments = edition_segments(values_row[1], values_row[2]) if index and highlighting and highlighting.enabled else None
            for column, (cell, text, width) in enumerate(zip(table.rows[index].cells, values_row, widths)):
                cell.width = width
                if segments is not None and column in (1, 2):
                    style = highlighting.deletions if column == 1 else highlighting.additions
                    cell.text = ""
                    for fragment, changed in segments[column - 1]:
                        run = cell.paragraphs[0].add_run(fragment)
                        if changed:
                            if style.color:
                                run.font.color.rgb = RGBColor.from_string(style.color)
                            if style.highlight:
                                marker = OxmlElement("w:highlight")
                                marker.set(qn("w:val"), style.highlight)
                                run._r.get_or_add_rPr().append(marker)
                            run.bold, run.italic, run.underline = style.bold, style.italic, style.underline
                            run.font.strike = getattr(style, "strike", False)
                else:
                    cell.text = text
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    for run in paragraph.runs:
                        run.font.name, run.font.size = "Times New Roman", Pt(11)
                        if index == 0:
                            run.bold = True

        doc.add_paragraph()
        for text in form.closing_text.splitlines():
            doc.add_paragraph(text)
        if form.signatures:
            heading = doc.add_paragraph("Подписи сторон")
            heading.runs[0].bold = True
            heading.paragraph_format.keep_with_next = True
            signatures = doc.add_table(rows=1, cols=len(form.signatures))
            signatures.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            # Counterparty left; the party issuing the protocol right.
            ordered = sorted(form.signatures.items(), key=lambda item: item[0] == drafting_party_id)
            for cell, (party_id, signature) in zip(signatures.rows[0].cells, ordered):
                cell.text = "\n".join(filter(None, [signature.role, signature.name])) or "Сторона: __________________"
                cell.add_paragraph(signature.position or "Должность: __________________")
                cell.add_paragraph("________________ / " + (signature.full_name or "________________") + " /")
                if form.mark_drafting_signature and party_id == drafting_party_id:
                    cell.add_paragraph("С учетом протокола разногласий")
                for paragraph in cell.paragraphs[:-1]:
                    paragraph.paragraph_format.keep_with_next = True

        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()
