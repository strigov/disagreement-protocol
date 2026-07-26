import zipfile
from io import BytesIO

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def build_simple_docx(paragraphs: list[str]) -> bytes:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_docx_with_insertion(
    before: str, inserted: str, after: str, author: str = "Editor"
) -> bytes:
    """Build a DOCX where `inserted` is wrapped in <w:ins>."""
    doc = DocxDocument()
    p = doc.add_paragraph(before)
    pelem = p._p
    ins = etree.SubElement(pelem, qn("w:ins"))
    ins.set(qn("w:id"), "1")
    ins.set(qn("w:author"), author)
    ins.set(qn("w:date"), "2026-01-01T00:00:00Z")
    r = etree.SubElement(ins, qn("w:r"))
    t = etree.SubElement(r, qn("w:t"))
    t.text = inserted
    t.set(qn("xml:space"), "preserve")
    if after:
        r2 = etree.SubElement(pelem, qn("w:r"))
        t2 = etree.SubElement(r2, qn("w:t"))
        t2.text = after
        t2.set(qn("xml:space"), "preserve")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_docx_with_table(rows: list[list[str]], leading_paragraph: str = "") -> bytes:
    """Build a DOCX with one table whose cells contain the given strings."""
    doc = DocxDocument()
    if leading_paragraph:
        doc.add_paragraph(leading_paragraph)
    n_rows = len(rows)
    n_cols = max(len(r) for r in rows) if rows else 0
    table = doc.add_table(rows=n_rows, cols=n_cols)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = val
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_docx_with_table_cell_replacement() -> bytes:
    """Build a table with one replacement and unchanged neighbouring cells."""
    doc = DocxDocument()
    doc.add_paragraph("Преамбула договора")
    doc.add_paragraph("1. Условия договора")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Период"
    table.cell(0, 1).text = "Сумма"
    table.cell(0, 2).text = "Итого"
    table.cell(1, 0).text = "2026 год"
    table.cell(1, 2).text = "192 500"

    replacement_paragraph = table.cell(1, 1).paragraphs[0]
    replacement_paragraph.text = ""
    deletion = etree.SubElement(replacement_paragraph._p, qn("w:del"))
    deletion.set(qn("w:id"), "1")
    deletion.set(qn("w:author"), "Юрист контрагента")
    deletion.set(qn("w:date"), "2026-01-01T00:00:00Z")
    deleted_run = etree.SubElement(deletion, qn("w:r"))
    deleted_text = etree.SubElement(deleted_run, qn("w:delText"))
    deleted_text.text = "1 200"
    deleted_text.set(qn("xml:space"), "preserve")

    insertion = etree.SubElement(replacement_paragraph._p, qn("w:ins"))
    insertion.set(qn("w:id"), "2")
    insertion.set(qn("w:author"), "Юрист контрагента")
    insertion.set(qn("w:date"), "2026-01-01T00:00:00Z")
    inserted_run = etree.SubElement(insertion, qn("w:r"))
    inserted_text = etree.SubElement(inserted_run, qn("w:t"))
    inserted_text.text = "1 100"
    inserted_text.set(qn("xml:space"), "preserve")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_docx_with_non_element_nodes() -> bytes:
    """Build a DOCX with XML comments and processing instructions in content."""
    doc = DocxDocument()
    body = doc._element.body

    doc.add_paragraph("До body")
    after_body_marker = doc.add_paragraph("После body")
    first_paragraph = body[0]
    body.insert(first_paragraph.getparent().index(first_paragraph) + 1, etree.Comment("между body blocks"))
    body.insert(first_paragraph.getparent().index(first_paragraph) + 2, etree.ProcessingInstruction("converter", "между body blocks"))

    run_paragraph = doc.add_paragraph()
    run = run_paragraph.add_run("До run")
    run_element = run._r
    run_element.append(etree.Comment("внутри run"))
    run_element.append(etree.ProcessingInstruction("converter", "внутри run"))
    after_run = etree.SubElement(run_element, qn("w:t"))
    after_run.text = " после run"
    after_run.set(qn("xml:space"), "preserve")

    table = doc.add_table(rows=1, cols=1)
    cell_paragraph = table.cell(0, 0).paragraphs[0]
    cell_run = cell_paragraph.add_run("До cell")
    cell_run_element = cell_run._r
    cell_run_element.append(etree.Comment("внутри cell"))
    cell_run_element.append(etree.ProcessingInstruction("converter", "внутри cell"))
    after_cell = etree.SubElement(cell_run_element, qn("w:t"))
    after_cell.text = " после cell"
    after_cell.set(qn("xml:space"), "preserve")

    assert after_body_marker._p.getparent() is body
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_docx_with_deletion(
    before: str, deleted: str, after: str, author: str = "Editor"
) -> bytes:
    """Build a DOCX where `deleted` is wrapped in <w:del>/<w:delText>."""
    doc = DocxDocument()
    p = doc.add_paragraph(before)
    pelem = p._p
    deln = etree.SubElement(pelem, qn("w:del"))
    deln.set(qn("w:id"), "1")
    deln.set(qn("w:author"), author)
    deln.set(qn("w:date"), "2026-01-01T00:00:00Z")
    r = etree.SubElement(deln, qn("w:r"))
    t = etree.SubElement(r, qn("w:delText"))
    t.text = deleted
    t.set(qn("xml:space"), "preserve")
    if after:
        r2 = etree.SubElement(pelem, qn("w:r"))
        t2 = etree.SubElement(r2, qn("w:t"))
        t2.text = after
        t2.set(qn("xml:space"), "preserve")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_docx_with_format_only_revisions() -> bytes:
    """Build a DOCX with one run and one paragraph formatting revision."""
    doc = DocxDocument()
    paragraph = doc.add_paragraph()
    ppr = paragraph._p.get_or_add_pPr()
    ppr_change = etree.SubElement(ppr, qn("w:pPrChange"))
    ppr_change.set(qn("w:id"), "11")
    ppr_change.set(qn("w:author"), "Автор абзаца")
    ppr_change.set(qn("w:date"), "2026-01-01T00:00:00Z")
    etree.SubElement(ppr_change, qn("w:pPr"))

    run = paragraph.add_run("Форматный текст")
    rpr = run._r.get_or_add_rPr()
    etree.SubElement(rpr, qn("w:b"))
    rpr_change = etree.SubElement(rpr, qn("w:rPrChange"))
    rpr_change.set(qn("w:id"), "12")
    rpr_change.set(qn("w:author"), "Автор run")
    rpr_change.set(qn("w:date"), "2026-01-01T00:00:00Z")
    etree.SubElement(rpr_change, qn("w:rPr"))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_docx_with_empty_format_only_run() -> bytes:
    """Build a DOCX with a format-only revision on an empty paragraph mark run."""
    doc = DocxDocument()
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    rpr = run._r.get_or_add_rPr()
    rpr_change = etree.SubElement(rpr, qn("w:rPrChange"))
    rpr_change.set(qn("w:id"), "13")
    rpr_change.set(qn("w:author"), "Автор paragraph mark")
    rpr_change.set(qn("w:date"), "2026-01-01T00:00:00Z")
    etree.SubElement(rpr_change, qn("w:rPr"))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_docx_with_many_run_format_only_revisions(change_count: int = 4) -> bytes:
    """Build one run containing several run-format change elements."""
    doc = DocxDocument()
    paragraph = doc.add_paragraph("Много форматных правок")
    rpr = paragraph.runs[0]._r.get_or_add_rPr()
    for index in range(change_count):
        rpr_change = etree.SubElement(rpr, qn("w:rPrChange"))
        rpr_change.set(qn("w:id"), str(index + 1))
        rpr_change.set(qn("w:author"), f"Автор {index + 1}")
        rpr_change.set(qn("w:date"), "2026-01-01T00:00:00Z")
        etree.SubElement(rpr_change, qn("w:rPr"))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_docx_with_manual_numbering(items: list[str]) -> bytes:
    """Items already include literal numbers like '1.', '1.1.', '2.'."""
    return build_simple_docx(items)


DEFAULT_NUMBERING_XML = f"""<w:numbering xmlns:w="{W_NS}">
  <w:abstractNum w:abstractNumId="0">
    <w:lvl w:ilvl="0">
      <w:start w:val="1"/>
      <w:numFmt w:val="decimal"/>
      <w:lvlText w:val="%1."/>
    </w:lvl>
    <w:lvl w:ilvl="1">
      <w:start w:val="1"/>
      <w:numFmt w:val="decimal"/>
      <w:lvlText w:val="%1.%2."/>
    </w:lvl>
    <w:lvl w:ilvl="2">
      <w:start w:val="1"/>
      <w:numFmt w:val="decimal"/>
      <w:lvlText w:val="%1.%2.%3."/>
    </w:lvl>
  </w:abstractNum>
  <w:num w:numId="1">
    <w:abstractNumId w:val="0"/>
  </w:num>
</w:numbering>"""


def build_docx_with_native_numbering(
    items: list[tuple[int | None, str]],
    numbering_xml: str = DEFAULT_NUMBERING_XML,
    styles_xml: str | None = None,
    num_id: int = 1,
    style_ids: list[str | None] | None = None,
) -> bytes:
    """Build a minimal DOCX package with real Word auto-numbering.

    `items` is a list of (ilvl, text); ilvl=None yields an unnumbered
    paragraph. `style_ids` (parallel to items) attaches w:pStyle without a
    direct numPr — for testing style-based numbering.
    """
    paragraphs = []
    for index, (ilvl, text) in enumerate(items):
        ppr_parts = []
        style_id = style_ids[index] if style_ids else None
        if style_id:
            ppr_parts.append(f'<w:pStyle w:val="{style_id}"/>')
        if ilvl is not None:
            ppr_parts.append(
                f'<w:numPr><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'
            )
        ppr = f"<w:pPr>{''.join(ppr_parts)}</w:pPr>" if ppr_parts else ""
        paragraphs.append(f'<w:p>{ppr}<w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>')
    document_xml = (
        f'<w:document xmlns:w="{W_NS}"><w:body>{"".join(paragraphs)}</w:body></w:document>'
    )

    content_types = (
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )

    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document_xml)
        z.writestr("word/numbering.xml", numbering_xml)
        if styles_xml is not None:
            z.writestr("word/styles.xml", styles_xml)
    return buf.getvalue()


def build_docx_with_unknown_numid(text: str, num_id: int = 999) -> bytes:
    """Build a DOCX where a paragraph references a numId that does not exist in numbering.xml."""
    doc = DocxDocument()
    p = doc.add_paragraph(text)
    pPr = p._p.get_or_add_pPr()
    numPr = etree.SubElement(pPr, qn("w:numPr"))
    ilvl = etree.SubElement(numPr, qn("w:ilvl"))
    ilvl.set(qn("w:val"), "0")
    numId = etree.SubElement(numPr, qn("w:numId"))
    numId.set(qn("w:val"), str(num_id))
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_autonumbering_contract() -> bytes:
    """A contract with manual numeric labels '1.' / '1.1.' / '2.', one INS revision in clause 1.1."""
    doc = DocxDocument()
    doc.add_paragraph(
        'ООО "Ромашка", в лице директора, именуемое в дальнейшем "Арендодатель", '
        'и АО "Василёк", именуемое в дальнейшем "Арендатор", заключили настоящий договор.'
    )
    doc.add_paragraph("1. Предмет договора")
    p = doc.add_paragraph("1.1. Исполнитель обязуется оказать услуги ")
    pelem = p._p
    ins = etree.SubElement(pelem, qn("w:ins"))
    ins.set(qn("w:id"), "1")
    ins.set(qn("w:author"), "Editor")
    ins.set(qn("w:date"), "2026-01-01T00:00:00Z")
    r = etree.SubElement(ins, qn("w:r"))
    t = etree.SubElement(r, qn("w:t"))
    t.text = "качественно и в срок"
    t.set(qn("xml:space"), "preserve")
    doc.add_paragraph("2. Цена договора")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_multipar_contract() -> bytes:
    """A clause that spans 2 paragraphs."""
    doc = DocxDocument()
    doc.add_paragraph(
        'ООО "А", именуемое в дальнейшем "Сторона1", и ООО "Б", именуемое в дальнейшем "Сторона2".'
    )
    doc.add_paragraph("1. Раздел один")
    p = doc.add_paragraph("Дополнительный текст пункта ")
    pelem = p._p
    ins = etree.SubElement(pelem, qn("w:ins"))
    ins.set(qn("w:id"), "1")
    ins.set(qn("w:author"), "X")
    ins.set(qn("w:date"), "2026-01-01T00:00:00Z")
    r = etree.SubElement(ins, qn("w:r"))
    t = etree.SubElement(r, qn("w:t"))
    t.text = "ВСТАВКА"
    t.set(qn("xml:space"), "preserve")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_new_clause_contract() -> bytes:
    """A whole paragraph wrapped in <w:ins>."""
    doc = DocxDocument()
    doc.add_paragraph(
        'ООО "А", именуемое в дальнейшем "Сторона1", и ООО "Б", именуемое в дальнейшем "Сторона2".'
    )
    doc.add_paragraph("1. Старый пункт")
    p = doc.add_paragraph("")
    pelem = p._p
    ins = etree.SubElement(pelem, qn("w:ins"))
    ins.set(qn("w:id"), "1")
    ins.set(qn("w:author"), "X")
    ins.set(qn("w:date"), "2026-01-01T00:00:00Z")
    r = etree.SubElement(ins, qn("w:r"))
    t = etree.SubElement(r, qn("w:t"))
    t.text = "2. Совершенно новый пункт"
    t.set(qn("xml:space"), "preserve")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_deleted_clause_contract() -> bytes:
    """A whole paragraph wrapped in <w:del>."""
    doc = DocxDocument()
    doc.add_paragraph(
        'ООО "А", именуемое в дальнейшем "Сторона1", и ООО "Б", именуемое в дальнейшем "Сторона2".'
    )
    doc.add_paragraph("1. Остающийся пункт")
    p = doc.add_paragraph("")
    pelem = p._p
    deln = etree.SubElement(pelem, qn("w:del"))
    deln.set(qn("w:id"), "1")
    deln.set(qn("w:author"), "X")
    deln.set(qn("w:date"), "2026-01-01T00:00:00Z")
    r = etree.SubElement(deln, qn("w:r"))
    t = etree.SubElement(r, qn("w:delText"))
    t.text = "2. Удаляемый пункт целиком"
    t.set(qn("xml:space"), "preserve")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_simple_contract_with_revision() -> bytes:
    return build_docx_with_insertion(
        'ООО "А", именуемое в дальнейшем "Сторона1", и ООО "Б", именуемое в дальнейшем "Сторона2". '
        "1. Услуги оказываются ",
        "вовремя",
        ".",
    )
