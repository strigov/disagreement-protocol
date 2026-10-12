from io import BytesIO

from docx import Document

from app.services.exporter import ProtocolExporter


def test_export_produces_docx_with_4col_table_and_headers():
    rows = [
        {
            "point_label": "1.",
            "party_1_text": "old",
            "party_2_text": "new",
            "agreed_text": "Редакция Арендатора",
        }
    ]

    data = ProtocolExporter().export(
        rows,
        drafting_role="Арендатор",
        counterparty_role="Арендодатель",
        contract_name="Договор",
    )

    doc = Document(BytesIO(data))
    assert any("Протокол разногласий" in p.text for p in doc.paragraphs)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.columns) == 4
    headers = [c.text for c in table.rows[0].cells]
    assert headers == [
        "Пункт договора",
        "Редакция Арендодателя",
        "Редакция Арендатора",
        "Согласованная редакция",
    ]
    body = [c.text for c in table.rows[1].cells]
    assert body == ["1.", "old", "new", "Редакция Арендатора"]


def test_export_falls_back_to_storona_labels_for_empty_roles():
    rows = [{"point_label": "1.", "party_1_text": "a", "party_2_text": "b", "agreed_text": "c"}]

    data = ProtocolExporter().export(rows, drafting_role="", counterparty_role="", contract_name=None)

    doc = Document(BytesIO(data))
    headers = [c.text for c in doc.tables[0].rows[0].cells]
    assert headers[1] == "Редакция Стороны 1"
    assert headers[2] == "Редакция Стороны 2"


def export_row(before, after, settings=None):
    from app.schemas.export_style import ChangeHighlighting
    data = ProtocolExporter().export(
        [{"point_label": "1.", "party_1_text": before, "party_2_text": after, "agreed_text": "Согласовано"}],
        "Покупатель", "Продавец", "Договор",
        highlighting=ChangeHighlighting.model_validate(settings) if settings is not None else None,
    )
    return Document(BytesIO(data)).tables[0].rows[1].cells


def test_highlights_only_changes_and_preserves_text():
    before, after = "Срок — 10 дней.\nЦена:\t100 рублей.", "Срок — 15 дней.\nЦена:\t100 рублей."
    cells = export_row(before, after, {"enabled": True})
    assert [cell.text for cell in cells] == ["1.", before, after, "Согласовано"]
    for cell, text, color in [(cells[1], "10", "FF0000"), (cells[2], "15", "00B050")]:
        marked = [r for p in cell.paragraphs for r in p.runs if r.font.color.rgb]
        assert [(r.text, str(r.font.color.rgb)) for r in marked] == [(text, color)]
    for cell in (cells[0], cells[3]):
        assert all(r.font.color.rgb is None for p in cell.paragraphs for r in p.runs)


def test_combines_all_styles_and_keeps_them_out_of_clean_version():
    style = {"color": "7030A0", "highlight": "cyan", "bold": True, "italic": True, "underline": True}
    settings = {"enabled": True, "deletions": {**style, "strike": True}, "additions": style}
    cells = export_row("старое", "новое", settings)
    for index in (1, 2):
        run = next(r for r in cells[index].paragraphs[0].runs if r.text)
        assert run.bold and run.italic and run.underline
        assert run.font.strike == (index == 1)
        assert str(run.font.color.rgb) == "7030A0"
        assert run._r.xpath("./w:rPr/w:highlight/@w:val") == ["cyan"]
    clean = export_row("старое", "новое", {**settings, "enabled": False})
    assert [c.text for c in cells] == [c.text for c in clean]
    for cell in clean:
        for run in cell.paragraphs[0].runs:
            assert not any((run.bold, run.italic, run.underline, run.font.strike, run.font.color.rgb, run.font.highlight_color))


def test_single_style_without_color():
    cells = export_row("старое", "новое", {
        "enabled": True, "deletions": {"color": None, "strike": True},
        "additions": {"color": None, "underline": True},
    })
    for index in (1, 2):
        run = next(r for r in cells[index].paragraphs[0].runs if r.text)
        assert run.font.color.rgb is None
        assert not run.bold and not run.italic
    assert cells[1].paragraphs[0].runs[-1].font.strike
    assert cells[2].paragraphs[0].runs[-1].underline


def test_empty_identical_and_repeated_editions():
    from app.services.export_diff import edition_segments
    for before, after in [("", "Новый пункт"), ("Удалённый пункт", ""), ("", ""), ("без изменений", "без изменений"), ("да, да; да.", "да, нет; да.")]:
        old, new = edition_segments(before, after)
        assert "".join(text for text, _ in old) == before
        assert "".join(text for text, _ in new) == after
        if before == after:
            assert not any(changed for _, changed in old + new)
    old, new = edition_segments("да " * 5000 + "нет", "да " * 5000 + "точно")
    assert [(text, changed) for text, changed in old if changed] == [("нет", True)]
    assert [(text, changed) for text, changed in new if changed] == [("точно", True)]


def test_service_texts_of_new_and_excluded_clauses_are_not_redlined():
    from app.services.export_diff import edition_segments
    from app.services.protocol_generator import ABSENT_EDITION, deletion_instruction
    original = "3.2. Поставщик вправе изменить цену, как указано в пункте 3.1."
    assert edition_segments(original, deletion_instruction("3.2."), "3.2.") == (
        [(original, True)], [(deletion_instruction("3.2."), False)],
    )
    assert edition_segments(ABSENT_EDITION, "Новая редакция", "Новый пункт 1") == (
        [(ABSENT_EDITION, False)], [("Новая редакция", True)],
    )
    # A manually rewritten instruction is ordinary wording again.
    old, new = edition_segments(original, "Исключить пункт 3.2 и пункт 3.1", "3.2.")
    assert any(changed for _, changed in new)


def test_redline_is_continuous_and_ignores_whitespace_only_edits():
    from app.services.export_diff import edition_segments
    old, new = edition_segments("Покупатель обязан оплатить товар.", "Продавец вправе потребовать оплату.")
    assert old == [("Покупатель обязан оплатить товар", True), (".", False)]
    assert new == [("Продавец вправе потребовать оплату", True), (".", False)]
    for before, after in [("срок 10 дней", "срок  10 дней"), ("срок 10 дней", "срок\n10 дней")]:
        old, new = edition_segments(before, after)
        assert not any(changed for _, changed in old + new)
        assert "".join(t for t, _ in old) == before and "".join(t for t, _ in new) == after
    old, _ = edition_segments("оплата в течение 10 дней", "оплата в течение 15 рабочих дней")
    assert old == [("оплата в течение ", False), ("10", True), (" дней", False)]
