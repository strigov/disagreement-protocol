from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.services.clause_codec import restore_clause, serialize_clause
from app.processing.types import Clause as ClauseData
from app.processing.validation_engine import ValidationEngine
from app.services.exporter import ProtocolExporter
from app.services.ingestion import run_pipeline
from app.services.protocol_generator import BuildContext, ProtocolRowBuilder


SAMPLES_DIR = Path(__file__).resolve().parents[3] / "samples"


def _sample_bytes(filename: str) -> bytes:
    return (SAMPLES_DIR / filename).read_bytes()


def _restore_clauses_from_persisted_paragraphs(result) -> list[ClauseData]:
    return [restore_clause(serialize_clause(clause, None, i))
            for i, clause in enumerate(result["clauses"])]


def test_sample_09_xml_comment_reaches_pipeline_result():
    result = run_pipeline(_sample_bytes("09_edge_xml_comment.docx"))

    assert result["clauses"]
    assert result["preflight_metrics"]["revisions_total"] == 1


def test_sample_03_counts_format_only_without_changed_clause():
    result = run_pipeline(_sample_bytes("03_complex_styles_annex_tables.docx"))

    assert result["preflight_metrics"]["format_only_count"] == 1
    format_only_clause_keys = {
        clause.clause_key
        for clause in result["clauses"]
        if any(
            revision.type == "FORMAT_ONLY"
            for paragraph in clause.paragraphs
            for revision in paragraph.revisions
        )
    }
    assert len(format_only_clause_keys) == 1
    assert format_only_clause_keys.isdisjoint(
        {clause.clause_key for clause in result["changed_clauses"]}
    )


def test_sample_03_extracts_two_parties_with_full_names_and_roles():
    result = run_pipeline(_sample_bytes("03_complex_styles_annex_tables.docx"))

    parties_by_role = {party.role: party for party in result["parties"]}
    assert set(parties_by_role) == {"Арендодатель", "Арендатор"}
    assert parties_by_role["Арендодатель"].name == (
        "Общество с ограниченной ответственностью «Гамма»"
    )
    assert parties_by_role["Арендатор"].name == "Акционерное общество «Дельта»"


def test_sample_03_table_replacement_uses_original_and_accepted_columns():
    result = run_pipeline(_sample_bytes("03_complex_styles_annex_tables.docx"))
    rows = ProtocolRowBuilder().build(
        result["changed_clauses"],
        BuildContext(drafting_party_role="Арендатор", counterparty_role="Арендодатель"),
    )

    table_rows = [row for row in rows if "2026 год" in row.party_1_text]
    assert len(table_rows) == 1
    row = table_rows[0]
    assert "2026 год | 1 200 | 192 500" in row.party_1_text
    assert "2026 год | 1 100 | 192 500" in row.party_2_text
    assert "1 2001 100" not in row.party_1_text + row.party_2_text
    assert "TABLE_CELL_CHANGE" in row.flags

    exported = ProtocolExporter().export(
        [
            {
                "point_label": row.point_label,
                "party_1_text": row.party_1_text,
                "party_2_text": row.party_2_text,
                "agreed_text": row.agreed_text,
            },
        ],
        "Арендатор",
        "Арендодатель",
        "Договор",
    )
    assert len(Document(BytesIO(exported)).tables) == 1


def test_sample_01_regex_numbering_has_no_per_paragraph_low_confidence_noise():
    result = run_pipeline(_sample_bytes("01_simple_manual_numbering.docx"))

    preflight_mean = result["preflight_metrics"]["numbering_confidence_mean"]
    assert preflight_mean == pytest.approx(0.4)
    restored_clauses = _restore_clauses_from_persisted_paragraphs(result)
    persisted_report = ValidationEngine().validate([], restored_clauses, [])
    assert persisted_report.numbering_confidence_mean == pytest.approx(preflight_mean)
    assert not any(
        diagnostic["code"] == "LOW_CONFIDENCE"
        for diagnostic in result["preflight_diagnostics"]
    )


def test_sample_02_preserves_deleted_numbers_and_uses_accepted_number_for_insertion():
    result = run_pipeline(_sample_bytes("02_medium_auto_numbering.docx"))
    rows = ProtocolRowBuilder().build(
        result["changed_clauses"],
        BuildContext(drafting_party_role="Арендатор", counterparty_role="Арендодатель"),
    )

    assert [row.point_label for row in rows] == ["1.1.", "1.3.", "2.1.", "2.2."]
    assert next(row for row in rows if "DELETED_CLAUSE" in row.flags).point_label == "2.2."
    assert not any(row.point_label == "2.3." for row in rows)


def test_sample_03_keeps_style_and_override_numbering():
    result = run_pipeline(_sample_bytes("03_complex_styles_annex_tables.docx"))

    assert [
        clause.label
        for clause in result["clauses"]
        if clause.label.startswith("Статья")
    ] == ["Статья I.", "Статья II.", "Статья III."]
    assert any(
        paragraph.rendered_number == "1.4.1."
        for clause in result["clauses"]
        for paragraph in clause.paragraphs
    )


def test_sample_10_uses_original_number_for_changed_clause_and_new_label_for_insertion():
    result = run_pipeline(_sample_bytes("10_edge_renumbering.docx"))
    rows = ProtocolRowBuilder().build(
        result["changed_clauses"],
        BuildContext(drafting_party_role="Поставщик", counterparty_role="Покупатель"),
    )

    assert [row.point_label for row in rows] == ["1.2.", "1.2."]
    assert "NEW_CLAUSE" in rows[0].flags
    changed_row = next(row for row in rows if "NEW_CLAUSE" not in row.flags)
    assert "14 календарных дней" in changed_row.party_1_text
    assert "7 календарных дней" in changed_row.party_2_text
