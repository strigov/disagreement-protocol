"""Public workflow regressions found in the 2026-08-31 review."""
import zipfile
from io import BytesIO

import pytest
from docx import Document as WordDocument
from sqlalchemy import event, inspect
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session, sessionmaker

from app.api.documents import get_analysis, update_parties
from app.api.protocols import add_row, build_protocol, delete_row, export_protocol, patch_row, swap_parties
from app.db.migrations import upgrade_schema
from app.models.clause import Clause
from app.models.document import Document
from app.models.enums import DocumentStatus
from app.models.party import Party
from app.models.protocol import Protocol, ProtocolRow
from app.processing.validation_engine import ValidationEngine
from app.schemas.party import PartiesUpdateIn
from app.schemas.protocol import ExportRequest, ProtocolRowCreate, ProtocolRowPatch
from app.services.clause_codec import serialize_clause
from app.services.ingestion import run_pipeline
from app.services.processing import process_document, recover_interrupted_documents
from tests.fixtures.builders import DEFAULT_NUMBERING_XML, build_docx_with_native_numbering

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def docx(body):
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("word/document.xml", f'<w:document xmlns:w="{W}"><w:body>{body}</w:body></w:document>')
    return buffer.getvalue()


def paragraph(text):
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def change(kind, text):
    tag = "delText" if kind == "del" else "t"
    return f'<w:{kind} w:id="3" w:author="Editor"><w:r><w:{tag}>{text}</w:{tag}></w:r></w:{kind}>'


def changed_paragraph(label, text):
    return f"<w:p><w:r><w:t>{label} </w:t></w:r>{change('ins', text)}</w:p>"


def table(text):
    return f"<w:tbl><w:tr><w:tc>{paragraph(text)}</w:tc></w:tr></w:tbl>"


def prepare(factory, data):
    result = run_pipeline(data)
    db = factory()
    doc = Document(original_filename="test.docx", storage_key="test", status=DocumentStatus.PARSED,
                   preflight_metrics=result["preflight_metrics"],
                   metadata_json={"pipeline_version": 2, "preflight_diagnostics": result["preflight_diagnostics"]})
    db.add(doc)
    db.flush()
    parties = [Party(document_id=doc.id, name="A", role="Поставщик", confidence=1),
               Party(document_id=doc.id, name="B", role="Покупатель", confidence=1)]
    db.add_all(parties)
    db.flush()
    doc.drafting_party_id = parties[1].id
    for i, clause in enumerate(result["clauses"]):
        db.add(serialize_clause(clause, doc.id, i))
    db.commit()
    return db, doc, parties, build_protocol(doc.id, db)


def ordered(db, protocol_id):
    return db.query(ProtocolRow).filter(ProtocolRow.protocol_id == protocol_id).order_by(ProtocolRow.sort_order).all()


def test_missing_rows_never_claim_covered_revisions():
    result = run_pipeline(docx(changed_paragraph("1.", "A")))
    report = ValidationEngine().validate([], result["clauses"], [])
    assert report.textual_revisions_total == 1
    assert report.covered_revisions == 0
    assert any(w["kind"] == "UNCOVERED_REVISIONS" for w in report.warnings)
    assert not any("не найдено" in w["message"] for w in report.warnings)


@pytest.mark.parametrize("direction,index", [("next", 0), ("previous", 1)])
def test_merge_preserves_order_coverage_and_append_position(isolated_database, direction, index):
    data = docx("".join(changed_paragraph(f"{i + 1}.", text) for i, text in enumerate("ABC")))
    db, doc, _, proto = prepare(isolated_database, data)
    try:
        original = ordered(db, proto.id)
        origins = [r.source_clause_key for r in original[:2]]
        merged = patch_row(proto.id, original[index].id, ProtocolRowPatch(merge_with=direction), db)
        assert merged.party_2_text == "1. A\n2. B"
        assert merged.point_label == "1., 2."
        assert db.get(ProtocolRow, merged.id).source_clause_keys == origins
        report = get_analysis(doc.id, db).validation_report
        assert (report.covered_revisions, report.textual_revisions_total) == (3, 3)
        new = add_row(proto.id, ProtocolRowCreate(), db)
        rows = ordered(db, proto.id)
        assert [r.id for r in rows] == [merged.id, original[2].id, new.id]
        assert [r.sort_order for r in rows] == [0, 1, 2]
        delete_row(proto.id, merged.id, db)
        assert doc.validation_report["covered_revisions"] == 1  # persisted before commit
        assert get_analysis(doc.id, db).validation_report.covered_revisions == 1
        patch_row(proto.id, original[2].id, ProtocolRowPatch(party_1_text="", party_2_text=""), db)
        assert get_analysis(doc.id, db).validation_report.covered_revisions == 0
        export_protocol(doc.id, ExportRequest(filename="test"), db)
        assert doc.validation_report["covered_revisions"] == 0
    finally:
        db.close()


def test_append_repairs_gaps_in_legacy_order(isolated_database):
    db, _, _, proto = prepare(isolated_database, docx("".join(changed_paragraph(f"{i}.", str(i)) for i in range(1, 4))))
    try:
        rows = ordered(db, proto.id)
        for row, position in zip(rows, [2, 4, 9]):
            row.sort_order = position
        db.commit()
        added = add_row(proto.id, ProtocolRowCreate(), db)
        after = ordered(db, proto.id)
        assert [r.id for r in after] == [*(r.id for r in rows), added.id]
        assert [r.sort_order for r in after] == list(range(4))
    finally:
        db.close()


def test_party_changes_sync_auto_agreed_and_preserve_explicit_edits(isolated_database):
    db, doc, parties, proto = prepare(isolated_database, docx("".join(changed_paragraph(f"{i}.", str(i)) for i in range(1, 4))))
    try:
        first, second, third = proto.rows
        patch_row(proto.id, first.id, ProtocolRowPatch(party_1_text="Ручная редакция текста"), db)
        patch_row(proto.id, second.id, ProtocolRowPatch(agreed_text="Ручное согласование"), db)
        # Even entering exactly the former generated label is an explicit edit.
        patch_row(proto.id, third.id, ProtocolRowPatch(agreed_text="Редакция Покупателя"), db)
        payload = [{"id": p.id, "name": p.name, "role": p.role} for p in parties]
        payload[0]["role"] = "Заказчик"
        update_parties(doc.id, PartiesUpdateIn(parties=payload, drafting_party_id=parties[0].id), db)
        rows = ordered(db, proto.id)
        assert [r.agreed_text for r in rows] == ["Редакция Заказчика", "Ручное согласование", "Редакция Покупателя"]
        assert rows[0].party_1_text == "Ручная редакция текста"
        exported = WordDocument(BytesIO(export_protocol(doc.id, ExportRequest(filename="test"), db).body))
        assert "Заказчика" in exported.tables[0].rows[0].cells[2].text
        assert exported.tables[0].rows[1].cells[3].text == "Редакция Заказчика"
        swap_parties(proto.id, db)
        assert rows[0].agreed_text == "Редакция Покупателя"
        assert rows[1].agreed_text == "Ручное согласование"
    finally:
        db.close()


def test_moves_are_versioned_counted_and_persisted(isolated_database):
    body = (f'<w:p><w:r><w:t>1. </w:t></w:r><w:moveFromRangeStart w:id="40" w:name="m"/>{change("moveFrom", "Аванс")}<w:moveFromRangeEnd w:id="40"/></w:p>'
            f'<w:p><w:r><w:t>2. </w:t></w:r><w:moveToRangeStart w:id="41" w:name="m"/>{change("moveTo", "Аванс")}<w:moveToRangeEnd w:id="41"/></w:p>')
    db, doc, _, proto = prepare(isolated_database, docx(body))
    try:
        assert [(r.party_1_text, r.party_2_text) for r in proto.rows] == [("1. Аванс", "1. "), ("2. ", "2. Аванс")]
        assert doc.preflight_metrics["revisions_total"] == 2
        assert doc.validation_report["covered_revisions"] == 2
        assert not any(d["code"] == "UNPAIRED_MOVE" for d in doc.metadata_json["preflight_diagnostics"])
        moved = WordDocument(BytesIO(export_protocol(doc.id, ExportRequest(filename="moves"), db).body))
        assert "Аванс" not in moved.tables[0].rows[1].cells[2].text
        assert "Аванс" in moved.tables[0].rows[2].cells[2].text
    finally:
        db.close()


def test_unpaired_move_is_reported():
    result = run_pipeline(docx(f"<w:p>{change('moveTo', 'Аванс')}</w:p>"))
    assert any(d["code"] == "UNPAIRED_MOVE" for d in result["preflight_diagnostics"])


@pytest.mark.parametrize("kind,deleted", [("ins", False), ("del", True), ("cellIns", False), ("cellDel", True)])
def test_structural_table_changes_survive_database_and_export(isolated_database, kind, deleted):
    properties = "tcPr" if kind.startswith("cell") else "trPr"
    mark = f'<w:{properties}><w:{kind} w:id="1" w:author="Editor"/></w:{properties}>'
    cell = (mark if properties == "tcPr" else "") + paragraph("Изменённый платёж")
    row = (mark if properties == "trPr" else "") + f"<w:tc>{cell}</w:tc>"
    body = paragraph("1. Платежи") + f"<w:tbl><w:tr>{row}</w:tr><w:tr><w:tc>{paragraph('Постоянный платёж')}</w:tc></w:tr></w:tbl>"
    db, doc, _, proto = prepare(isolated_database, docx(body))
    try:
        result = proto.rows[0]
        assert ("Изменённый платёж" in result.party_1_text) is deleted
        assert ("Изменённый платёж" in result.party_2_text) is not deleted
        assert "Постоянный платёж" in result.party_1_text and "Постоянный платёж" in result.party_2_text
        assert doc.preflight_metrics["revisions_total"] == 1
        assert doc.validation_report["covered_revisions"] == 1
        exported = WordDocument(BytesIO(export_protocol(doc.id, ExportRequest(filename="table"), db).body))
        assert ("Изменённый платёж" in exported.tables[0].rows[1].cells[2].text) is not deleted
    finally:
        db.close()


def test_unsupported_cell_merge_is_not_silently_accepted():
    with pytest.raises(ValueError, match="cellMerge"):
        run_pipeline(docx('<w:tbl><w:tr><w:tc><w:tcPr><w:cellMerge w:id="1"/></w:tcPr></w:tc></w:tr></w:tbl>'))


def test_mixed_block_order_survives_database_and_export(isolated_database):
    body = (paragraph("1. До таблицы") + table("ТАБЛИЦА 1")
            + f"<w:p><w:r><w:t>Между </w:t></w:r>{change('ins', 'НОВОЕ')}</w:p>"
            + table("ТАБЛИЦА 2") + paragraph("После таблиц"))
    db, doc, _, proto = prepare(isolated_database, docx(body))
    try:
        expected = "1. До таблицы\nТАБЛИЦА 1\nМежду НОВОЕ\nТАБЛИЦА 2\nПосле таблиц"
        assert proto.rows[0].party_2_text == expected
        exported = WordDocument(BytesIO(export_protocol(doc.id, ExportRequest(filename="order"), db).body))
        assert exported.tables[0].rows[1].cells[2].text == expected
    finally:
        db.close()


@pytest.mark.parametrize("restart,last", [("0", "2.2."), ("1", "2.1."), (None, "2.1."), ("9", "2.1.")])
def test_lvl_restart_rules(restart, last):
    numbering = DEFAULT_NUMBERING_XML
    if restart is not None:
        numbering = numbering.replace('<w:lvl w:ilvl="1">', f'<w:lvl w:ilvl="1"><w:lvlRestart w:val="{restart}"/>')
    data = build_docx_with_native_numbering([(0, "Раздел"), (1, "Пункт"), (0, "Раздел"), (1, "Пункт")], numbering_xml=numbering)
    result = run_pipeline(data)
    assert [p.rendered_number for c in result["clauses"] for p in c.paragraphs] == ["1.", "1.1.", "2.", last]


def test_failed_result_persistence_marks_failed_after_rollback(isolated_database, monkeypatch):
    with isolated_database() as db:
        doc = Document(original_filename="test.docx", storage_key="key", status=DocumentStatus.UPLOADED)
        db.add(doc)
        db.commit()
        document_id = doc.id
        engine = db.get_bind()

    class FailingSession(Session):
        pass

    @event.listens_for(FailingSession, "before_flush")
    def fail_result(session, *_):
        if any(isinstance(obj, Clause) for obj in session.new):
            raise OperationalError("INSERT clauses", {}, Exception("simulated failure"))

    bad_factory = sessionmaker(bind=engine, class_=FailingSession, autoflush=False, expire_on_commit=False)
    monkeypatch.setattr("app.services.processing.SessionLocal", bad_factory)
    monkeypatch.setattr("app.services.processing.FileStorage", lambda: type("Storage", (), {"get": lambda _, key: docx(changed_paragraph('1.', 'A'))})())
    process_document(str(document_id))
    with isolated_database() as db:
        failed = db.get(Document, document_id)
        assert failed.status == DocumentStatus.FAILED
        assert "simulated failure" in failed.error_message
        assert db.query(Clause).count() == 0


def test_recovery_and_processing_are_idempotent(isolated_database, monkeypatch):
    with isolated_database() as db:
        docs = [Document(original_filename="test.docx", storage_key="key", status=status) for status in DocumentStatus]
        db.add_all(docs)
        db.commit()
        before = {doc.id: doc.status for doc in docs}
        assert recover_interrupted_documents(db) == 2
        assert recover_interrupted_documents(db) == 0
        for doc in docs:
            if before[doc.id] not in (DocumentStatus.UPLOADED, DocumentStatus.PARSING):
                assert doc.status == before[doc.id]
        fresh = Document(original_filename="test.docx", storage_key="key", status=DocumentStatus.UPLOADED)
        db.add(fresh)
        db.commit()
        document_id = fresh.id
    monkeypatch.setattr("app.services.processing.FileStorage", lambda: type("Storage", (), {"get": lambda _, key: docx(changed_paragraph('1.', 'A'))})())
    process_document(str(document_id))
    process_document(str(document_id))
    with isolated_database() as db:
        assert db.get(Document, document_id).status == DocumentStatus.PARSED
        assert db.query(Clause).filter(Clause.document_id == document_id).count() == 1


def test_additive_migration_preserves_old_rows_and_manual_edits(isolated_database):
    db, doc, _, proto = prepare(isolated_database, docx(changed_paragraph("1.", "A")))
    row = db.get(ProtocolRow, proto.rows[0].id)
    row.agreed_text = "Ручная редакция до обновления"
    row.user_edited = True
    db.commit()
    row_id, engine = row.id, db.get_bind()
    db.close()
    with engine.begin() as connection:
        connection.exec_driver_sql("ALTER TABLE protocol_rows DROP COLUMN source_clause_keys")
        connection.exec_driver_sql("ALTER TABLE protocol_rows DROP COLUMN agreed_text_edited")
    upgrade_schema(engine)
    upgrade_schema(engine)
    assert {c["name"] for c in inspect(engine).get_columns("protocol_rows")} >= {"source_clause_keys", "agreed_text_edited"}
    with isolated_database() as db:
        row = db.get(ProtocolRow, row_id)
        assert row.agreed_text_edited is None
        assert row.agreed_text == "Ручная редакция до обновления"
        assert row.user_edited is True
        swap_parties(proto.id, db)
        assert row.agreed_text == "Ручная редакция до обновления"
        assert db.query(Protocol).count() == 1 and db.query(Clause).count() == 1


def test_old_block_order_is_recovered_from_original_without_losing_text(isolated_database, monkeypatch):
    source = docx(paragraph("1. До") + table("ТАБЛИЦА") + changed_paragraph("После", "НОВОЕ"))
    db, doc, _, proto = prepare(isolated_database, source)
    try:
        before = proto.rows[0].party_2_text
        for model in db.query(Clause).all():
            model.paragraphs_json = [{**p, "block_index": None} for p in model.paragraphs_json]
            model.tables_json = [{k: v for k, v in t.items() if k != "block_index"} for t in model.tables_json]
        db.commit()
        monkeypatch.setattr("app.api.protocols.FileStorage.get", lambda _, key: source)
        rebuilt = build_protocol(doc.id, db)
        assert rebuilt.rows[0].party_2_text == before == "1. До\nТАБЛИЦА\nПосле НОВОЕ"
        assert not any(w.kind == "SOURCE_ORDER_UNAVAILABLE" for w in get_analysis(doc.id, db).validation_report.warnings)
    finally:
        db.close()


@pytest.mark.parametrize("stage", ["storage_init", "storage_read", "pipeline"])
def test_processing_failure_at_every_early_stage_is_terminal(isolated_database, monkeypatch, stage):
    with isolated_database() as db:
        doc = Document(original_filename="test.docx", storage_key="key", status=DocumentStatus.UPLOADED)
        db.add(doc)
        db.commit()
        document_id = doc.id
    class Storage:
        def __init__(self):
            if stage == "storage_init":
                raise OSError("storage init failure")
        def get(self, key):
            if stage == "storage_read":
                raise OSError("storage read failure")
            return b"not a ZIP"
    monkeypatch.setattr("app.services.processing.FileStorage", Storage)
    process_document(str(document_id))
    with isolated_database() as db:
        failed = db.get(Document, document_id)
        assert failed.status == DocumentStatus.FAILED and failed.error_message
        assert db.query(Clause).count() == 0
