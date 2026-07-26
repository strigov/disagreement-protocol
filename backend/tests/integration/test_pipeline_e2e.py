from io import BytesIO

from docx import Document

from app.processing.changed_clause_detector import ChangedClauseDetector
from app.services.exporter import ProtocolExporter
from app.services.ingestion import run_pipeline
from app.services.protocol_generator import BuildContext, ProtocolRowBuilder
from tests.fixtures.builders import (
    build_autonumbering_contract,
    build_deleted_clause_contract,
    build_docx_with_format_only_revisions,
    build_multipar_contract,
    build_new_clause_contract,
    build_simple_contract_with_revision,
    build_docx_with_table_cell_replacement,
    build_docx_with_native_numbering,
)


def _full_run(data: bytes, drafting_role: str, counterparty_role: str):
    pipeline = run_pipeline(data)
    changed = ChangedClauseDetector().detect(pipeline["clauses"])
    rows = ProtocolRowBuilder().build(
        changed,
        BuildContext(
            drafting_party_role=drafting_role,
            counterparty_role=counterparty_role,
        ),
    )
    docx_bytes = ProtocolExporter().export(
        [
            {
                "point_label": r.point_label,
                "party_1_text": r.party_1_text,
                "party_2_text": r.party_2_text,
                "agreed_text": r.agreed_text,
            }
            for r in rows
        ],
        drafting_role,
        counterparty_role,
        "Договор",
    )
    return rows, docx_bytes


def test_simple_contract_yields_one_row_and_valid_docx():
    rows, data = _full_run(build_simple_contract_with_revision(), "Сторона2", "Сторона1")
    assert len(rows) == 1
    doc = Document(BytesIO(data))
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 2


def test_autonumbering_contract_pipeline():
    rows, _ = _full_run(build_autonumbering_contract(), "Арендатор", "Арендодатель")
    assert any(r.point_label.startswith("1.1") for r in rows)


def test_multipar_contract_keeps_continuation():
    rows, _ = _full_run(build_multipar_contract(), "Сторона2", "Сторона1")
    assert any("ВСТАВКА" in r.party_2_text for r in rows)


def test_new_clause_contract_produces_new_pungkt_label():
    rows, _ = _full_run(build_new_clause_contract(), "Сторона2", "Сторона1")
    new_rows = [r for r in rows if r.point_label.startswith("Новый пункт")]
    assert len(new_rows) >= 1
    assert new_rows[0].party_1_text == "—"


def test_deleted_clause_contract_produces_exclude_text():
    rows, _ = _full_run(build_deleted_clause_contract(), "Сторона2", "Сторона1")
    excl = [r for r in rows if r.party_2_text.startswith("Исключить пункт")]
    assert len(excl) >= 1


def test_format_only_revision_counts_without_creating_changed_clause():
    result = run_pipeline(build_docx_with_format_only_revisions())

    assert result["preflight_metrics"]["format_only_count"] == 2
    format_only_clause_keys = {
        clause.clause_key
        for clause in result["clauses"]
        if any(
            revision.type == "FORMAT_ONLY"
            for paragraph in clause.paragraphs
            for revision in paragraph.revisions
        )
    }
    changed_clause_keys = {clause.clause_key for clause in result["changed_clauses"]}
    assert format_only_clause_keys
    assert format_only_clause_keys.isdisjoint(changed_clause_keys)


def test_table_cell_replacement_survives_pipeline_row_and_export():
    result = run_pipeline(build_docx_with_table_cell_replacement())
    rows = ProtocolRowBuilder().build(
        result["changed_clauses"],
        BuildContext(drafting_party_role="Арендатор", counterparty_role="Арендодатель"),
    )

    table_rows = [row for row in rows if "2026 год" in row.party_1_text]
    assert len(table_rows) == 1
    row = table_rows[0]
    assert "2026 год | 1 200 | 192 500" in row.party_1_text
    assert "2026 год | 1 100 | 192 500" in row.party_2_text
    assert "1 2001 100" not in row.party_1_text + row.party_2_text
    assert "TABLE_CELL_CHANGE" in row.flags
    assert result["preflight_metrics"]["revisions_total"] == 2

    exported = ProtocolExporter().export(
        [
            {
                "point_label": row.point_label,
                "party_1_text": row.party_1_text,
                "party_2_text": row.party_2_text,
                "agreed_text": row.agreed_text,
            },
        ],
        "Арендатор",
        "Арендодатель",
        "Договор",
    )
    exported_doc = Document(BytesIO(exported))
    assert len(exported_doc.tables) == 1


def test_inserted_native_clause_keeps_original_numbers_and_new_clause_semantics():
    data = build_docx_with_native_numbering(
        [(0, "Раздел"), (1, "Исходный подпункт"), (1, "Вставленный подпункт"), (1, "Следующий подпункт")],
        content_revisions=[None, None, "INS", None],
        paragraph_mark_revisions=[None, None, "INS", None],
    )

    result = run_pipeline(data)
    rows = ProtocolRowBuilder().build(
        result["changed_clauses"],
        BuildContext(drafting_party_role="Сторона2", counterparty_role="Сторона1"),
    )

    assert [paragraph.rendered_number for clause in result["clauses"] for paragraph in clause.paragraphs] == [
        "1.",
        "1.1.",
        None,
        "1.2.",
    ]
    new_rows = [row for row in rows if row.point_label.startswith("Новый пункт")]
    assert len(new_rows) == 1
    assert new_rows[0].party_1_text == "—"
    assert new_rows[0].party_2_text == "Вставленный подпункт"
    assert [row.point_label for row in rows if "NEW_CLAUSE" not in row.flags] == []


def test_deleted_native_clause_keeps_original_number():
    data = build_docx_with_native_numbering(
        [(0, "Первый пункт"), (0, "Удалённый пункт"), (0, "Следующий пункт")],
        content_revisions=[None, "DEL", None],
        paragraph_mark_revisions=[None, "DEL", None],
    )

    result = run_pipeline(data)
    rows = ProtocolRowBuilder().build(
        result["changed_clauses"],
        BuildContext(drafting_party_role="Сторона2", counterparty_role="Сторона1"),
    )

    deleted_rows = [row for row in rows if "DELETED_CLAUSE" in row.flags]
    assert len(deleted_rows) == 1
    assert deleted_rows[0].point_label == "2."
    assert deleted_rows[0].party_2_text == "Исключить пункт 2."
