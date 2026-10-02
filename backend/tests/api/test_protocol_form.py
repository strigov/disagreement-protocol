"""Form persistence, defaults, party ownership and the public DOCX boundary."""
import uuid
from io import BytesIO

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, inspect, text

from app.db.migrations import upgrade_schema
from app.main import app
from app.models.document import Document
from app.models.enums import DocumentStatus
from app.models.party import Party
from app.models.protocol import Protocol, ProtocolRow


@pytest.fixture
def prepared(isolated_database):
    with isolated_database() as db:
        doc = Document(original_filename="example.docx", storage_key="example", contract_name="поставки № 12", status=DocumentStatus.BUILT)
        db.add(doc)
        db.flush()
        parties = [Party(document_id=doc.id, name="ООО Поставщик", role="Поставщик", confidence=1),
                   Party(document_id=doc.id, name="ООО Покупатель", role="Покупатель", confidence=1)]
        db.add_all(parties)
        db.flush()
        doc.drafting_party_id = parties[1].id
        protocol = Protocol(document_id=doc.id)
        db.add(protocol)
        db.flush()
        db.add(ProtocolRow(protocol_id=protocol.id, sort_order=0, point_label="1.1", party_1_text="Исходный текст",
                           party_2_text="Новая редакция", agreed_text="Редакция Покупателя", agreed_text_edited=False))
        db.commit()
        return str(doc.id), str(protocol.id), [str(p.id) for p in parties]


def test_form_survives_reload_and_export_without_changing_rows(prepared, isolated_database):
    doc, protocol, parties = prepared
    with TestClient(app, base_url="http://localhost") as client:
        before = client.get(f"/api/protocols/{protocol}/rows").json()
        form = before["form"]
        assert form["date"] == "" and form["place"] == ""
        assert form["contract_reference"] == "поставки № 12"
        assert "ООО Покупатель" in form["opening_text"]
        patch = {"title": "Протокол разногласий № 3", "date": "02.10.2026", "place": "г. Москва",
                 "opening_text": "Наша шапка\nВторая строка", "closing_text": "Наши условия",
                 "original_title": "Текст поставщика", "agreed_title": "Итог",
                 "signatures": {parties[1]: {"full_name": "Иванов И. И.", "position": "Директор"}}}
        response = client.patch(f"/api/protocols/{protocol}/form", json=patch)
        assert response.status_code == 200, response.text
        reloaded = client.get(f"/api/protocols/{protocol}/rows").json()
        assert reloaded["form"] == response.json()
        assert reloaded["rows"] == before["rows"]
        assert client.get(f"/api/protocols?document_id={doc}").json()[0]["form"] == response.json()
        exported = client.post(f"/api/documents/{doc}/protocol/export", json={"filename": "Протокол"})
        assert exported.status_code == 200
        word = WordDocument(BytesIO(exported.content))
        text_content = "\n".join(p.text for p in word.paragraphs)
        assert "Наша шапка\nВторая строка" in text_content
        assert "Наши условия" in text_content and "02.10.2026" in text_content
        assert word.tables[0].cell(0, 1).text == "Текст поставщика"
        assert word.tables[0].cell(1, 2).text == "Новая редакция"
        assert "Иванов И. И." in word.tables[1].cell(0, 1).text
        assert "С учетом протокола разногласий" in word.tables[1].cell(0, 1).text
        assert "С учетом протокола разногласий" not in word.tables[1].cell(0, 0).text
        assert word.tables[0].rows[0]._tr.xpath("./w:trPr/w:tblHeader")
        assert word.sections[0].page_width > word.sections[0].page_height
        client.patch(f"/api/protocols/{protocol}/form", json={"place": "Казань"})
    with isolated_database() as db:
        assert db.get(Document, uuid.UUID(doc)).status == DocumentStatus.BUILT
        assert db.get(Protocol, uuid.UUID(protocol)).form_json["place"] == "Казань"


def test_swap_keeps_signature_owners_and_manual_titles_but_updates_defaults(prepared):
    doc, protocol, parties = prepared
    with TestClient(app, base_url="http://localhost") as client:
        client.patch(f"/api/protocols/{protocol}/form", json={"original_title": "Своя колонка", "signatures": {parties[1]: {"full_name": "Петров"}}})
        form = client.post(f"/api/protocols/{protocol}/swap-parties").json()["form"]
        assert form["original_title"] == "Своя колонка"
        assert form["proposed_title"] == "Редакция Поставщика"
        assert form["signatures"][parties[1]]["full_name"] == "Петров"
        word = WordDocument(BytesIO(client.post(f"/api/documents/{doc}/protocol/export", json={"filename": "test"}).content))
        assert "Петров" in word.tables[1].cell(0, 0).text
        assert "С учетом протокола разногласий" not in word.tables[1].cell(0, 0).text
        assert "ООО Поставщик" in word.tables[1].cell(0, 1).text
        assert "С учетом протокола разногласий" in word.tables[1].cell(0, 1).text
        restored = client.patch(f"/api/protocols/{protocol}/form", json={"original_title": None}).json()
        assert restored["original_title"] == "Редакция Покупателя"
        client.patch(f"/api/protocols/{protocol}/form", json={"mark_drafting_signature": False, "closing_text": ""})
        word = WordDocument(BytesIO(client.post(f"/api/documents/{doc}/protocol/export", json={"filename": "test"}).content))
        assert not any("С учетом протокола разногласий" in cell.text for row in word.tables[1].rows for cell in row.cells)
        assert not any("Во всем остальном" in p.text for p in word.paragraphs)


@pytest.mark.parametrize("patch", [
    {"title": "  "}, {"original_title": ""}, {"place": "x" * 501}, {"closing_text": "x" * 20001},
    {"opening_text": "bad\x00text"}, {"place": 123}, {"unknown": "no"},
    {"signatures": {str(uuid.uuid4()): {"full_name": "Чужая сторона"}}},
    {"signatures": {"not-a-uuid": {"name": "X"}}},
])
def test_invalid_form_is_rejected_atomically(prepared, patch):
    _, protocol, _ = prepared
    with TestClient(app, base_url="http://localhost") as client:
        before = client.get(f"/api/protocols/{protocol}/rows").json()
        response = client.patch(f"/api/protocols/{protocol}/form", json=patch)
        assert response.status_code == 422, response.text
        assert client.get(f"/api/protocols/{protocol}/rows").json() == before


def test_rebuild_retains_form(prepared):
    doc, protocol, _ = prepared
    with TestClient(app, base_url="http://localhost") as client:
        saved = client.patch(f"/api/protocols/{protocol}/form", json={"opening_text": "Своя шапка"}).json()
        response = client.post(f"/api/documents/{doc}/protocol")
        assert response.status_code == 201
        assert response.json()["form"] == saved


def test_additive_form_migration_preserves_existing_data(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'legacy.db'}")
    upgrade_schema(engine)
    with engine.begin() as db:
        db.exec_driver_sql("ALTER TABLE protocols DROP COLUMN form_json")
        db.execute(text("INSERT INTO protocols(id, document_id) VALUES (:id, :doc)"), {"id": "a" * 32, "doc": "b" * 32})
    upgrade_schema(engine)
    upgrade_schema(engine)
    with engine.connect() as db:
        assert "form_json" in {col["name"] for col in inspect(db).get_columns("protocols")}
        assert db.execute(text("SELECT id, document_id, form_json FROM protocols")).one() == ("a" * 32, "b" * 32, None)
    engine.dispose()


def test_standard_opening_tracks_signature_names_while_custom_text_is_preserved(prepared):
    _, protocol, parties = prepared
    with TestClient(app, base_url="http://localhost") as client:
        response = client.patch(f"/api/protocols/{protocol}/form", json={"signatures": {parties[1]: {"name": "ООО Новое имя"}}})
        assert "ООО Новое имя" in response.json()["opening_text"]
        client.patch(f"/api/protocols/{protocol}/form", json={"opening_text": "Наш собственный текст"})
        response = client.patch(f"/api/protocols/{protocol}/form", json={"signatures": {parties[1]: {"name": "ООО Еще одно имя"}}})
        assert response.json()["opening_text"] == "Наш собственный текст"
        restored = client.patch(f"/api/protocols/{protocol}/form", json={"opening_text": None}).json()
        assert "ООО Еще одно имя" in restored["opening_text"]
