"""Exercise the shipped executable with an isolated database, never the developer DB."""
import argparse
from contextlib import contextmanager
from io import BytesIO
import json
import os
from pathlib import Path
import queue
import secrets
import subprocess
import tempfile
import threading
import time

from docx import Document
import httpx
from lxml import etree


@contextmanager
def engine(executable, data_dir):
    token = secrets.token_hex(32)
    env = {**os.environ, "DISPROT_DESKTOP_TOKEN": token, "DATABASE_URL": "sqlite:///must-not-be-used.db"}
    process = subprocess.Popen(
        [str(executable), "--data-dir", str(data_dir)], cwd=data_dir.parent,
        env=env, stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
        text=True, encoding="utf-8",
    )
    lines = queue.Queue()
    threading.Thread(target=lambda: [lines.put(line) for line in process.stdout], daemon=True).start()
    try:
        line = lines.get(timeout=60)
        assert line.startswith("DISPROT_READY="), line
        port = json.loads(line.partition("=")[2])["port"]
        with httpx.Client(base_url=f"http://127.0.0.1:{port}", timeout=30) as client:
            deadline = time.monotonic() + 30
            while True:
                try:
                    if client.get("/health", headers={"x-disprot-token": token}).status_code == 200:
                        break
                except httpx.TransportError:
                    pass
                assert process.poll() is None and time.monotonic() < deadline, "Engine startup failed"
                time.sleep(0.1)
            assert client.get("/health").status_code == 403
            client.headers.update({"x-disprot-token": token, "origin": "disprot://app"})
            yield client
        # EOF also simulates losing the Electron parent process.
        process.stdin.close()
        assert process.wait(timeout=10) == 0, "Engine did not stop cleanly"
    finally:
        if process.poll() is None:
            process.kill()
        process.wait(timeout=10)


def synthetic_contract():
    document = Document()
    document.add_paragraph("1. Услуги оказываются в срок.")
    p = document.add_paragraph()
    qn = lambda name: "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + name
    ins = etree.SubElement(p._p, qn("ins"), {qn("id"): "1", qn("author"): "Тест"})
    etree.SubElement(etree.SubElement(ins, qn("r")), qn("t")).text = "2. Новый пункт договора."
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def checked(response):
    response.raise_for_status()
    return response.json()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("executable", type=Path)
    parser.add_argument("--contract", type=Path)
    args = parser.parse_args()
    executable = args.executable.resolve()
    source = args.contract.read_bytes() if args.contract else synthetic_contract()
    with tempfile.TemporaryDirectory(prefix="disprot-smoke-") as temp:
        root = Path(temp)
        data_dir = root / "Данные с пробелами"
        with engine(executable, data_dir) as client:
            document_id = checked(client.post("/api/documents", files={"file": ("Договор.docx", source)}))["id"]
            deadline = time.monotonic() + 30
            while True:
                analysis = checked(client.get(f"/api/documents/{document_id}/analysis"))
                if analysis["status"] not in ("uploaded", "parsing"):
                    break
                assert time.monotonic() < deadline, "Parsing timed out"
                time.sleep(0.1)
            assert analysis["status"] == "parsed", analysis
            parties = [{"id": None, "name": "ООО Альфа", "role": "Заказчик"}, {"id": None, "name": "ООО Бета", "role": "Исполнитель"}]
            analysis = checked(client.patch(f"/api/documents/{document_id}/parties", json={"parties": parties, "drafting_party_id": None}))
            parties = [{key: p[key] for key in ("id", "name", "role")} for p in analysis["parties"]]
            checked(client.patch(f"/api/documents/{document_id}/parties", json={"parties": parties, "drafting_party_id": parties[0]["id"]}))
            protocol = checked(client.post(f"/api/documents/{document_id}/protocol"))
            assert protocol["rows"]
            if not args.contract:
                assert protocol["rows"][0]["point_label"].rstrip(".") == "2"
            row = protocol["rows"][0]
            checked(client.patch(f'/api/protocols/{protocol["id"]}/rows/{row["id"]}', json={"agreed_text": "Сохранённая ручная редакция"}))
            exported = client.post(f"/api/documents/{document_id}/protocol/export", json={"filename": "Протокол"})
            exported.raise_for_status()
            output = Document(BytesIO(exported.content))
            assert len(output.tables[0].columns) == 4
            assert "Сохранённая ручная редакция" in " ".join(c.text for r in output.tables[0].rows for c in r.cells)
            print(f"Upload, parse, edit and DOCX export: OK ({len(protocol['rows'])} rows)")
        assert (data_dir / "disprot.db").exists()
        assert not (root / "must-not-be-used.db").exists()
        with engine(executable, data_dir) as client:
            restored = checked(client.get(f'/api/protocols/{protocol["id"]}/rows'))
            assert restored["rows"][0]["agreed_text"] == "Сохранённая ручная редакция"
        print("Authentication, Unicode paths, restart persistence and shutdown: OK")


if __name__ == "__main__":
    main()
